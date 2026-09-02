"""Universal Document Loader supporting ANY file format.

Rich formats (PDF, DOCX, PPTX, XLSX, CSV, JSON, HTML, images with OCR) are
converted to a markdown-flavoured plain text so the downstream chunkers can
treat every document the same way. Anything else is read as text when it is
decodable, and reported as unsupported otherwise.

Metadata (article_id, product_area, last_updated) comes from either:
  * YAML frontmatter at the top of a text/markdown file, or
  * a sidecar ``<filename>.meta.json`` next to a binary file.
"""

import json
import csv
import re
from html.parser import HTMLParser
from pathlib import Path
from datetime import datetime

import yaml

DATA_DIR = Path("data")

META_SUFFIX = ".meta.json"

# Extensions that we store as plain text with YAML frontmatter.
TEXT_FRONTMATTER_EXTS = {".md", ".markdown", ".txt", ".text", ".rst"}

# Image types we try to OCR (requires pillow + pytesseract, both optional).
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif", ".webp"}


class ExtractionError(Exception):
    """Raised when no usable text can be extracted from a file."""


class _HTMLTextExtractor(HTMLParser):
    """Minimal HTML → text converter using only the standard library."""

    _BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "header",
              "footer", "h1", "h2", "h3", "h4", "h5", "h6", "table", "pre",
              "blockquote", "ul", "ol"}
    _SKIP = {"script", "style", "noscript", "head"}

    def __init__(self):
        super().__init__()
        self._parts = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        if tag in ("h1", "h2", "h3"):
            self._parts.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self._parts.append("\n- ")
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self._parts.append(data)

    def text(self):
        raw = "".join(self._parts)
        raw = re.sub(r"[ \t]+", " ", raw)
        raw = re.sub(r"\n\s*\n\s*\n+", "\n\n", raw)
        return raw.strip()


def _rows_to_markdown(rows):
    """Render a list of row-lists as a markdown table."""
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [[str(c).replace("\n", " ").strip() for c in r] + [""] * (width - len(r)) for r in rows]
    header = " | ".join(rows[0])
    sep = " | ".join(["---"] * width)
    body = "\n".join("| " + " | ".join(r) + " |" for r in rows[1:])
    return f"| {header} |\n| {sep} |\n{body}".rstrip()


def _looks_like_text(data: bytes) -> bool:
    """Heuristic: is this byte blob plausibly a text file?"""
    if not data:
        return False
    sample = data[:4096]
    if b"\x00" in sample:
        return False
    printable = sum(1 for b in sample if 32 <= b < 127 or b in (9, 10, 13) or b >= 128)
    return printable / len(sample) > 0.9


def _read_text_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_text_from_file(file_path: Path) -> str:
    """Extract clean text content from any file format.

    Raises ExtractionError if nothing usable could be extracted.
    """
    ext = file_path.suffix.lower()

    try:
        # 1. PDF
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(file_path))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                txt = page.extract_text() or ""
                if txt.strip():
                    pages_text.append(f"## Page {idx + 1}\n{txt.strip()}")
            text = "\n\n".join(pages_text)
            if not text.strip():
                raise ExtractionError(
                    "PDF contains no extractable text (it may be scanned images)."
                )
            return text

        # 2. Word
        if ext == ".docx":
            import docx
            doc = docx.Document(str(file_path))
            elements = []
            for p in doc.paragraphs:
                if p.text.strip():
                    style = p.style.name if p.style else ""
                    if "Heading" in style or style == "Title":
                        elements.append(f"## {p.text.strip()}")
                    else:
                        elements.append(p.text.strip())
            for table in doc.tables:
                md = _rows_to_markdown([[c.text for c in row.cells] for row in table.rows])
                if md:
                    elements.append(md)
            return "\n\n".join(elements)

        # 3. PowerPoint (optional dependency)
        if ext == ".pptx":
            try:
                from pptx import Presentation
            except ImportError:
                raise ExtractionError(
                    "PowerPoint support needs `pip install python-pptx`."
                )
            prs = Presentation(str(file_path))
            slides = []
            for idx, slide in enumerate(prs.slides, start=1):
                lines = []
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        lines.append(shape.text_frame.text.strip())
                    if getattr(shape, "has_table", False) and shape.has_table:
                        md = _rows_to_markdown(
                            [[c.text for c in row.cells] for row in shape.table.rows]
                        )
                        if md:
                            lines.append(md)
                if lines:
                    slides.append(f"## Slide {idx}\n" + "\n\n".join(lines))
            return "\n\n".join(slides)

        # 4. Excel
        if ext in (".xlsx", ".xlsm", ".xls"):
            if ext == ".xls":
                raise ExtractionError(
                    "Legacy .xls is not supported — re-save the workbook as .xlsx."
                )
            import openpyxl
            wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
            sheet_texts = []
            for sheetname in wb.sheetnames:
                rows = [[c if c is not None else "" for c in r]
                        for r in wb[sheetname].iter_rows(values_only=True)]
                md = _rows_to_markdown(rows)
                if md:
                    sheet_texts.append(f"## Sheet: {sheetname}\n\n{md}")
            return "\n\n".join(sheet_texts)

        # 5. CSV / TSV
        if ext in (".csv", ".tsv"):
            delimiter = "\t" if ext == ".tsv" else ","
            with open(file_path, "r", encoding="utf-8", errors="ignore", newline="") as f:
                rows = list(csv.reader(f, delimiter=delimiter))
            return _rows_to_markdown(rows)

        # 6. JSON
        if ext == ".json":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return json.dumps(json.load(f), indent=2, ensure_ascii=False)

        # 7. HTML
        if ext in (".html", ".htm", ".xhtml"):
            parser = _HTMLTextExtractor()
            parser.feed(file_path.read_text(encoding="utf-8", errors="ignore"))
            return parser.text()

        # 8. Images → OCR when available
        if ext in IMAGE_EXTS:
            try:
                from PIL import Image
                import pytesseract
            except ImportError:
                raise ExtractionError(
                    "Image OCR needs `pip install pillow pytesseract` and the "
                    "tesseract binary. Upload a text-based file instead."
                )
            text = pytesseract.image_to_string(Image.open(file_path))
            if not text.strip():
                raise ExtractionError("OCR found no readable text in the image.")
            return text.strip()

        # 9. Everything else: Markdown, TXT, source code, configs, logs, XML, YAML…
        data = file_path.read_bytes()
        if not _looks_like_text(data):
            raise ExtractionError(
                f"'{file_path.suffix or 'no extension'}' is a binary format with no "
                "text extractor. Supported: PDF, DOCX, PPTX, XLSX, CSV, JSON, HTML, "
                "images (OCR) and any plain-text file."
            )
        return _read_text_bytes(data).strip()

    except ExtractionError:
        raise
    except Exception as e:  # parser failure of any kind
        raise ExtractionError(f"Could not read {file_path.name}: {e}") from e


def meta_path_for(file_path: Path) -> Path:
    return file_path.with_name(file_path.name + META_SUFFIX)


def _default_metadata(file_path: Path, article_content: str) -> dict:
    lines = article_content.split("\n")
    title = lines[0].lstrip("# ").strip() if lines and lines[0].startswith("#") else file_path.stem
    return {
        "article_id": file_path.stem.upper().replace(" ", "-"),
        "title": title,
        "product_area": "Custom Upload",
        "last_updated": datetime.now().strftime("%Y-%m-%d"),
    }


def load_articles():
    """Load and parse all documents in the data directory regardless of file type."""
    documents = []
    DATA_DIR.mkdir(exist_ok=True)

    for file_path in sorted(DATA_DIR.glob("*")):
        if file_path.name.startswith(".") or file_path.is_dir():
            continue
        if file_path.name.endswith(META_SUFFIX):
            continue

        try:
            raw_content = extract_text_from_file(file_path)
        except ExtractionError as e:
            print(f"  [loader] skipping {file_path.name}: {e}")
            continue
        if not raw_content or not raw_content.strip():
            continue

        # YAML frontmatter (text/markdown files)
        parts = raw_content.split("---", 2)
        if len(parts) == 3 and parts[0].strip() == "":
            article_content = parts[2].strip()
            try:
                metadata = yaml.safe_load(parts[1]) or {}
            except Exception:
                metadata = {}
            if not isinstance(metadata, dict):
                metadata = {}
        else:
            article_content = raw_content.strip()
            metadata = _default_metadata(file_path, article_content)

        # Sidecar metadata (binary files) overrides defaults.
        sidecar = meta_path_for(file_path)
        if sidecar.exists():
            try:
                metadata.update(json.loads(sidecar.read_text(encoding="utf-8")))
            except Exception:
                pass

        metadata.setdefault("source_file", file_path.name)
        if not metadata.get("article_id"):
            metadata["article_id"] = file_path.stem.upper().replace(" ", "-")
        if not metadata.get("product_area"):
            metadata["product_area"] = "General"
        if not metadata.get("last_updated"):
            metadata["last_updated"] = datetime.now().strftime("%Y-%m-%d")
        metadata["file_type"] = file_path.suffix.lower().lstrip(".") or "txt"

        documents.append({"content": article_content, "metadata": metadata})

    return documents


if __name__ == "__main__":
    docs = load_articles()
    print(f"Loaded {len(docs)} documents:")
    for d in docs:
        print(f" - {d['metadata']['article_id']}: {d['metadata']['source_file']} ({len(d['content'])} chars)")
